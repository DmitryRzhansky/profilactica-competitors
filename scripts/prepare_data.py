# -*- coding: utf-8 -*-
import csv
import json
import shutil
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(r"c:\Users\User\Downloads\profilactica-competitors")
DATA = ROOT / "assets" / "data"
DATA.mkdir(parents=True, exist_ok=True)


def read_csv(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


arch = read_csv(ROOT / "profilactica_final_architecture.csv")
geo = read_csv(ROOT / "profilactica_geo_pages.csv")
master = read_csv(ROOT / "profilactica_geo_master_lists.csv")

print("arch keys", list(arch[0].keys()) if arch else None)
print("arch rows", len(arch))
print("priority", Counter(r["priority"] for r in arch))
print("silos", Counter(r["silo"] for r in arch))
print(
    "geo_policy nonempty",
    sum(1 for r in arch if (r.get("geo_policy") or "").strip() not in ("", "Нет")),
)

print("geo pages", len(geo))
print("by service", Counter(r["service_name"] for r in geo))
print("by type", Counter(r["geo_type"] for r in geo))
print("by tier", Counter(r["tier"] for r in geo))
matrix = defaultdict(lambda: Counter())
for r in geo:
    matrix[r["service_name"]][r["geo_type"]] += 1
for name, counts in matrix.items():
    print(name, dict(counts), "total", sum(counts.values()))

print("master", len(master), Counter(r["geo_type"] for r in master), Counter(r["tier"] for r in master))

wb = openpyxl.load_workbook(ROOT / "result.xlsx", read_only=True, data_only=True)
pages = []
for i, row in enumerate(wb["_summary"].iter_rows(values_only=True)):
    if i == 0:
        continue
    if row and row[1] is not None:
        pages.append(int(row[1]))
print("competitor domains", len(pages), "urls", sum(pages))
import statistics

print("median", statistics.median(pages), "mean", round(statistics.mean(pages), 1))


def dump(sheet, needles, n=25):
    ws = wb[sheet]
    print("\n===", sheet)
    hits = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        url = (row[4] or "") if row and len(row) > 4 else ""
        title = (row[2] or "") if row and len(row) > 2 else ""
        blob = (str(url) + " " + str(title)).lower()
        if any(x.lower() in blob for x in needles):
            hits.append((url, str(title)[:110]))
    for h in hits[:n]:
        print(h[0], "|", h[1])
    print("hits", len(hits))


dump("alcomed.ru", ["/vyvod-iz-zapoya", "/ceny", "/vrachi", "/stacionar", "лечение алкоголизма"])
dump("klinika-korsakov.ru", ["psihiatr", "психиатр", "stacionar", "uslugi", "depress", "психиатр"])
dump("narkolog24.clinic", ["/ceny", "tarif", "narkolog-na-dom", "/price", "/service/"])
dump("ne-zavisimost.ru", ["semya", "rodstven", "motivac", "posle-lechen", "sozavis"])
dump("verimed.ru", ["narkologicheskaya-pomosh", "/uslugi/narkologiya"])
dump("alcoclinic.ru", ["/uslugi/kodirovanie-ot-alkogolizma"])
wb.close()

# Copy CSVs into assets/data
shutil.copy2(ROOT / "profilactica_final_architecture.csv", DATA / "architecture.csv")
shutil.copy2(ROOT / "profilactica_geo_pages.csv", DATA / "geo-pages.csv")
shutil.copy2(ROOT / "profilactica_geo_master_lists.csv", DATA / "geo-master.csv")
print("copied csvs")
